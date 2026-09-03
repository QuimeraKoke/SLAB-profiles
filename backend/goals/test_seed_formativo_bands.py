"""Guards for `seed_formativo_bands` and the band metadata the editor needs.

The failure mode here is an inverted scale. If the direction is read wrong, the
fastest players get flagged as the worst and the alert list looks perfectly
plausible — nobody recomputes a seven-band scale by hand. The club's own sheet
has that typo in one block, which is why the direction is inferred from the
numbers and not from the text.
"""
from __future__ import annotations

from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.test import TestCase

from api.alert_rules import build_rule_meta
from core.management.commands.backfill_cohorts import BRACKETS
from core.models import Bracket, Category, Club, Department
from exams.models import ExamTemplate
from goals.management.commands import seed_formativo_bands as cmd
from goals.models import AlertRule, AlertRuleKind

BANDAS = cmd.BANDAS


class BuildRangesTests(TestCase):
    """`build_ranges` alone — the piece that decides the orientation."""

    def test_umbrales_que_suben_son_mas_es_mejor(self):
        # RM Back Squat, Sub 13: 68,3 → 142,7.
        ranges, mas_es_mejor = cmd.build_ranges(
            [68.3, 86.0, 100.4, 115.6, 125.1, 142.7])
        self.assertTrue(mas_es_mejor)
        self.assertEqual(ranges[0], {"label": "Muy Deficiente",
                                     "color": ranges[0]["color"], "max": 68.3})
        self.assertEqual(ranges[-1]["min"], 142.7)
        self.assertNotIn("max", ranges[-1], "Elite queda abierta hacia arriba")

    def test_umbrales_que_bajan_son_menos_es_mejor(self):
        # T10, Sub 20: 1,8 → 1,53. Un tiempo mejora al bajar.
        ranges, mas_es_mejor = cmd.build_ranges([1.8, 1.72, 1.66, 1.59, 1.56, 1.53])
        self.assertFalse(mas_es_mejor)
        self.assertEqual(ranges[0]["min"], 1.8)
        self.assertNotIn("max", ranges[0], "Muy Deficiente abierta hacia arriba")
        self.assertEqual(ranges[-1]["max"], 1.53)
        self.assertNotIn("min", ranges[-1])

    def test_las_bandas_intermedias_no_dejan_huecos_ni_se_solapan(self):
        for umbrales in ([68.3, 86.0, 100.4, 115.6, 125.1, 142.7],
                         [1.8, 1.72, 1.66, 1.59, 1.56, 1.53]):
            ranges, _ = cmd.build_ranges(umbrales)
            self.assertEqual(len(ranges), len(BANDAS))
            for a, b in zip(ranges, ranges[1:]):
                # El borde compartido tiene que ser el MISMO número: un hueco
                # deja valores sin banda y un solapamiento los pone en dos.
                bordes = {a.get("min"), a.get("max")} & {b.get("min"), b.get("max")}
                self.assertTrue(bordes - {None},
                                f"{a['label']} y {b['label']} no comparten borde")

    def test_todas_las_bandas_llevan_etiqueta_y_al_menos_un_limite(self):
        # Es lo que valida `AlertRule.clean()`; si no, la regla no guarda.
        ranges, _ = cmd.build_ranges([1.8, 1.72, 1.66, 1.59, 1.56, 1.53])
        for b in ranges:
            self.assertTrue(b["label"])
            self.assertTrue("min" in b or "max" in b)


CAB = ["FECHA", "JUGADOR", "FECHA DE NACIMIENTO", "EDAD", "CATEGORÍA", "POSICIÓN"]


def escribir_formato(ruta: Path, bloques: list[dict]) -> None:
    """Reproduce la forma del club: grupo en la columna A dentro del bloque."""
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    # Las dos hojas tienen que existir o el comando aborta.
    for hoja in cmd.SHEETS:
        ws = book.create_sheet(hoja)
        fila = 4
        for bloque in [b for b in bloques if b["hoja"] == hoja]:
            ws.cell(row=fila, column=3, value=bloque["test"])
            ws.cell(row=fila, column=4, value=bloque["unidad"])
            ws.cell(row=fila, column=5, value="Rango")
            for i, label in enumerate(BANDAS):
                r = fila + 1 + i
                ws.cell(row=r, column=3, value=label)
                if i < len(bloque["umbrales"]):
                    ws.cell(row=r, column=4, value=bloque["umbrales"][i])
                if i == 0 and bloque.get("texto"):
                    ws.cell(row=r, column=5, value=bloque["texto"])
            # La etiqueta del grupo vive DENTRO del bloque, como en el original.
            ws.cell(row=fila + 3, column=1, value=bloque["grupo"])
            fila += len(BANDAS) + 4
    book.save(ruta)


class SeedBandsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.club = Club.objects.create(name="Universidad de Chile")
        cls.dept = Department.objects.create(club=cls.club, name="Físico",
                                             slug="fisico")
        for order, (code, name, age) in enumerate(BRACKETS):
            Bracket.objects.create(code=code, name=name, age=age, order=order)
        cls.brackets = {b.age: b for b in Bracket.objects.all()}
        # Serie 2013 = Sub 13 en 2026; Serie 2006 no existe, el mayor es SUB-20.
        cls.s2013 = Category.objects.create(club=cls.club, name="Serie 2013",
                                            cohort_year=2013)
        cls.s2012 = Category.objects.create(club=cls.club, name="Serie 2012",
                                            cohort_year=2012)
        cls.sub20 = Category.objects.create(club=cls.club, name="SUB-20")
        cls.sub20.team_seasons.create(season=2026, bracket=cls.brackets[20])
        cls.fuerza = ExamTemplate.objects.create(
            department=cls.dept, name="Fuerza", slug="fuerza",
            config_schema={"fields": [
                {"key": "rm_estimado", "label": "1RM", "type": "number",
                 "unit": "kg", "chart_type": "line"},
                {"key": "peso_corporal", "label": "Peso", "type": "number",
                 "unit": "kg"},
            ]})
        cls.fuerza.applicable_categories.set([cls.s2013, cls.s2012, cls.sub20])
        cls.carreras = ExamTemplate.objects.create(
            department=cls.dept, name="Carreras", slug="carreras",
            config_schema={"fields": [
                {"key": "t10_best", "label": "T10", "type": "calculated",
                 "unit": "s", "chart_type": "line"},
            ]})
        cls.carreras.applicable_categories.set([cls.s2013, cls.s2012, cls.sub20])

    def correr(self, bloques, *args):
        with TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "eval.xlsx"
            escribir_formato(ruta, bloques)
            salida = StringIO()
            call_command("seed_formativo_bands", "--file", str(ruta),
                         "--club", self.club.name, "--season", "2026",
                         *args, stdout=salida)
            return salida.getvalue()

    FUERZA_13 = {
        "hoja": cmd.SHEETS[1], "grupo": "SUB 13 - SUB 14",
        "test": "RM BACK SQUAT", "unidad": "kg", "texto": "<68,3",
        "umbrales": [68.3, 86.0, 100.4, 115.6, 125.1, 142.7, ">142,7"],
    }
    T10_20 = {
        "hoja": cmd.SHEETS[0], "grupo": "18-20",
        "test": "T10", "unidad": "Tiempo (seg)", "texto": "<1,8",   # el typo
        "umbrales": [1.8, 1.72, 1.66, 1.59, 1.56, 1.53, ">1,53"],
    }

    def test_crea_una_regla_por_categoria_del_grupo(self):
        # "SUB 13 - SUB 14" cubre dos peldaños → dos categorías.
        self.correr([self.FUERZA_13], "--commit")
        reglas = AlertRule.objects.filter(template=self.fuerza,
                                          field_key="rm_estimado")
        self.assertEqual({r.category.name for r in reglas},
                         {"Serie 2013", "Serie 2012"})
        r = reglas.get(category=self.s2013)
        self.assertEqual(r.kind, AlertRuleKind.BAND)
        self.assertEqual(r.config["ranges"][0]["max"], 68.3)
        self.assertEqual(r.config["trigger_labels"], ["Muy Deficiente"])

    def test_el_texto_invertido_no_manda_sobre_los_numeros(self):
        """El bloque T10 de `18-20` dice `<1,8` y los umbrales bajan.

        Si ganara el texto, "Muy Deficiente" sería *menos* de 1,8 s — o sea que
        los más rápidos quedarían marcados como los peores, y la lista de
        alertas se vería perfectamente razonable.
        """
        salida = self.correr([self.T10_20], "--commit")
        r = AlertRule.objects.get(template=self.carreras, field_key="t10_best",
                                  category=self.sub20)
        peor = r.config["ranges"][0]
        self.assertEqual(peor["min"], 1.8, "el peor T10 es el más LENTO")
        self.assertNotIn("max", peor)
        self.assertIn("Texto del rango contradice", salida)

    def test_el_plantel_mayor_se_resuelve_por_TeamSeason(self):
        # SUB-20 no tiene cohorte propia: su peldaño sale de la TeamSeason.
        self.correr([self.T10_20], "--commit")
        self.assertTrue(AlertRule.objects.filter(category=self.sub20).exists())

    def test_no_toca_una_regla_existente_sin_overwrite(self):
        self.correr([self.FUERZA_13], "--commit")
        r = AlertRule.objects.get(category=self.s2013, field_key="rm_estimado")
        r.config = {**r.config, "ranges": [{"label": "Mía", "max": 1}]}
        r.save()
        salida = self.correr([self.FUERZA_13], "--commit")
        r.refresh_from_db()
        self.assertEqual(r.config["ranges"][0]["label"], "Mía")
        self.assertIn("ya existían", salida)

    def test_con_overwrite_re_siembra_los_numeros_y_respeta_los_disparadores(self):
        """Se re-siembran los umbrales, no la decisión de qué dispara.

        Elegir qué banda alerta es del cuerpo técnico; los números son del
        club. Pisar `trigger_labels` borraría trabajo ajeno.
        """
        self.correr([self.FUERZA_13], "--commit")
        r = AlertRule.objects.get(category=self.s2013, field_key="rm_estimado")
        r.config = {**r.config, "trigger_labels": ["Deficiente", "Regular"]}
        r.save()
        self.correr([self.FUERZA_13], "--commit", "--overwrite")
        r.refresh_from_db()
        self.assertEqual(r.config["ranges"][0]["max"], 68.3)
        self.assertEqual(r.config["trigger_labels"], ["Deficiente", "Regular"])

    def test_salta_la_categoria_a_la_que_la_plantilla_no_aplica(self):
        # El club define bandas de YOYO para Sub 11–12 aunque `resistencia`
        # arranca en Sub 13. Ampliar la aplicabilidad es otra decisión.
        self.fuerza.applicable_categories.remove(self.s2012)
        salida = self.correr([self.FUERZA_13], "--commit")
        self.assertFalse(AlertRule.objects.filter(category=self.s2012).exists())
        self.assertIn("no aplica", salida)

    def test_sin_commit_no_escribe(self):
        self.correr([self.FUERZA_13])
        self.assertEqual(AlertRule.objects.count(), 0)

    def test_un_bloque_sin_mapeo_se_reporta(self):
        raro = {**self.FUERZA_13, "test": "TEST QUE NO EXISTE"}
        salida = self.correr([raro], "--commit")
        self.assertEqual(AlertRule.objects.count(), 0)
        self.assertIn("sin mapeo", salida.lower())


class BandMetaTests(TestCase):
    """El editor necesita los NÚMEROS, no sólo las etiquetas."""

    @classmethod
    def setUpTestData(cls):
        cls.club = Club.objects.create(name="Universidad de Chile")
        cls.dept = Department.objects.create(club=cls.club, name="Físico",
                                             slug="fisico")
        cls.cat = Category.objects.create(club=cls.club, name="Serie 2013",
                                          cohort_year=2013)
        cls.tpl = ExamTemplate.objects.create(
            department=cls.dept, name="Carreras", slug="carreras",
            config_schema={"fields": [
                {"key": "t10_1", "label": "T10 intento 1", "type": "number",
                 "unit": "s"},
                {"key": "t10_best", "label": "T10 mejor", "type": "calculated",
                 "unit": "s", "chart_type": "line"},
                {"key": "cod_asimetria", "label": "Asimetría",
                 "type": "calculated", "unit": "%",
                 "reference_ranges": [
                     {"label": "Simétrico", "min": -10, "max": 10,
                      "color": "#16a34a"},
                     {"label": "Asimétrico", "min": 10, "color": "#dc2626"},
                 ]},
            ]})
        cls.tpl.applicable_categories.set([cls.cat])

    def _campos(self):
        meta = build_rule_meta(self.cat)
        tpl = next(t for t in meta["templates"] if t["slug"] == "carreras")
        return {f["key"]: f for f in tpl["band_fields"]}, tpl["band_fields"]

    def test_todo_campo_numerico_puede_llevar_bandas(self):
        """El filtro anterior era `if f.get("reference_ranges")`.

        Con umbrales por regla eso escondía justo los campos que los necesitan:
        las cinco plantillas del club casi no tienen bandas en el campo, así
        que las 66 reglas sembradas apuntaban a campos que el formulario no
        ofrecía.
        """
        campos, _ = self._campos()
        self.assertIn("t10_best", campos)
        self.assertIn("t10_1", campos, "nada se esconde")

    def test_manda_los_numeros_de_las_bandas_del_campo(self):
        campos, _ = self._campos()
        defaults = campos["cod_asimetria"]["default_ranges"]
        self.assertEqual(len(defaults), 2)
        self.assertEqual(defaults[0]["min"], -10)
        self.assertEqual(defaults[0]["max"], 10)

    def test_un_campo_sin_bandas_propias_manda_lista_vacia(self):
        campos, _ = self._campos()
        self.assertEqual(campos["t10_best"]["default_ranges"], [])
        self.assertEqual(campos["t10_best"]["bands"], [])

    def test_los_campos_para_leer_van_primero(self):
        # Abrir la lista a todos los campos la volvió ruidosa: `carreras` pasó
        # de 1 opción a 21, tres de ellas los intentos de un mismo sprint.
        _, orden = self._campos()
        claves = [f["key"] for f in orden]
        self.assertIn("t10_best", [f["key"] for f in orden if f["featured"]])
        self.assertIn("t10_1", [f["key"] for f in orden if not f["featured"]])
        self.assertLess(claves.index("t10_best"), claves.index("t10_1"))

    def test_sigue_mandando_las_etiquetas_para_los_chips(self):
        campos, _ = self._campos()
        self.assertEqual(campos["cod_asimetria"]["bands"],
                         ["Simétrico", "Asimétrico"])
