"""Carga de archivo para Fatiga Central (CFF), con la planilla del club.

Encabezados que trae el archivo:

    Jugador · I1 (Hz) · I2 (Hz) · I3 (Hz) · CFF mean · CFF basal ·
    Δ% vs basal · Var % intra-sesión · PR · EA · N° alertas

De esas once, **sólo cinco se mapean**. Las cuatro calculadas y el conteo de
alertas se ignoran a propósito, y estos tests fijan esa decisión con la razón:
mapear un campo `calculated` es un no-op —`compute_result_data` lo pisa— y las
celdas de la planilla traen artefactos, cosa que el importador histórico ya
había documentado.

Dos de estos tests cubren bugs reales que aparecieron al habilitar la primera
plantilla con `column_mapping`. La carga masiva por UI nunca se había ejercido
completa, así que ninguno se había disparado:

  * el endpoint dejaba `recorded_at` naive y el `post_save` de goals reventaba
    con `localtime() cannot be applied to a naive datetime`;
  * el commit no era atómico, así que ese 500 dejaba filas a medias.
"""
from __future__ import annotations

import io
from unittest import mock

import openpyxl
from django.test import Client, TestCase

from api.auth import issue_token
from core.models import Category, Club, Department, Player, StaffMembership
from exams.models import ExamResult, ExamTemplate

HEADERS = [
    "Jugador", "I1 (Hz)", "I2 (Hz)", "I3 (Hz)", "CFF mean", "CFF basal",
    "Δ% vs basal", "Var % intra-sesión", "PR", "EA", "N° alertas",
]


def workbook(rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for j, h in enumerate(HEADERS, 1):
        ws.cell(row=1, column=j, value=h)
    for i, row in enumerate(rows, start=2):
        for j, h in enumerate(HEADERS, 1):
            ws.cell(row=i, column=j, value=row.get(h))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FatigaCentralUploadTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from django.core.management import call_command

        cls.club = Club.objects.create(name="Universidad de Chile")
        cls.dept = Department.objects.create(
            club=cls.club, name="Psicosocial", slug="psicosocial")
        cls.cat = Category.objects.create(
            club=cls.club, name="Primer Equipo", is_senior=True)
        cls.cat.departments.add(cls.dept)
        call_command(
            "seed_fatiga_central", "--create-if-missing",
            "--club", "Universidad de Chile", verbosity=0,
        )
        cls.template = ExamTemplate.objects.get(
            slug="fatiga_central", is_active_version=True)
        # Las reglas de alerta importan para estos tests, no son decorado: el
        # bug del datetime naive vivía en el post_save que las evalúa, así que
        # sin ellas el camino no se ejerce y el test pasa sin probar nada.
        # Verificado quitando `make_aware`: con reglas falla, sin reglas no.
        call_command(
            "seed_fatiga_alert_rules", "--club", "Universidad de Chile",
            verbosity=0,
        )
        cls.player = Player.objects.create(
            category=cls.cat, first_name="Charles", last_name="Aranguiz")
        cls.other = Player.objects.create(
            category=cls.cat, first_name="Elías", last_name="Rojas")

    def setUp(self):
        from django.contrib.auth import get_user_model

        user = get_user_model().objects.create_superuser(
            username="root", password="x")
        m = StaffMembership.objects.create(
            user=user, club=self.club, all_departments=True)
        # La membresía necesita la categoría: el endpoint filtra por
        # `scope_categories`, y un superusuario CON membresía queda sujeto al
        # alcance de esa membresía (el bypass total es sólo sin membresía).
        m.categories.add(self.cat)
        self.client = Client(HTTP_AUTHORIZATION=f"Bearer {issue_token(user)[0]}")

    def _post(self, rows, *, dry_run=True, recorded_at="2026-08-20T12:00:00"):
        return self.client.post("/api/results/bulk", {
            "file": io.BytesIO(workbook(rows)),
            "template_id": str(self.template.id),
            "category_id": str(self.cat.id),
            "recorded_at": recorded_at,
            "dry_run": "true" if dry_run else "false",
        })

    def _row(self, name, **over):
        row = {
            "Jugador": name,
            "I1 (Hz)": 41.2, "I2 (Hz)": 40.8, "I3 (Hz)": 41.6,
            "PR": 7, "EA": 8,
        }
        row.update(over)
        return row

    # ── la configuración de la plantilla ────────────────────────────────
    def test_the_template_offers_bulk_ingest(self):
        self.assertIn("bulk_ingest", self.template.input_config["input_modes"])

    def test_only_the_five_real_inputs_are_mapped(self):
        """Las calculadas y el conteo de alertas quedan fuera, a propósito."""
        mapped = set(
            self.template.input_config["column_mapping"]["field_map"].keys())
        self.assertEqual(
            mapped, {"I1 (Hz)", "I2 (Hz)", "I3 (Hz)", "PR", "EA"})
        for excluded in ("CFF mean", "CFF basal", "Δ% vs basal",
                         "Var % intra-sesión", "N° alertas"):
            self.assertNotIn(excluded, mapped, excluded)

    # ── el camino completo ──────────────────────────────────────────────
    def test_a_file_with_the_clubs_headers_matches_its_players(self):
        r = self._post([self._row("Charles Aranguiz"), self._row("Elías Rojas")])
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body["matched_players"], 2)
        self.assertEqual(body["unmatched"], [])

    def test_the_computed_columns_are_recomputed_not_trusted(self):
        """La razón nº1 para no mapearlas: son `calculated` y el motor las pisa.

        El archivo trae basura deliberada donde van los cálculos —incluida una
        FECHA en "CFF mean", que es exactamente el artefacto que documentó el
        importador histórico— y nada de eso debe llegar al resultado.
        """
        r = self._post([self._row(
            "Charles Aranguiz",
            **{"CFF mean": "2026-01-01", "CFF basal": 99.9,
               "Δ% vs basal": -77.7, "Var % intra-sesión": 88.8},
        )])
        data = r.json()["matched"][0]["result_data"]

        self.assertEqual(data["cff_mean"], 41.2)          # (41.2+40.8+41.6)/3
        self.assertEqual(data["var_intra_pct"], 1.94)     # (41.6-40.8)/41.2
        self.assertNotEqual(data["cff_basal"], 99.9)
        self.assertNotEqual(data["delta_basal_pct"], -77.7)

    def test_the_alert_count_does_not_reach_the_result(self):
        # SLAB cuenta sus propias alertas desde las bandas del template. Traer
        # el conteo de la planilla sería una segunda fuente de lo mismo.
        r = self._post([self._row("Charles Aranguiz", **{"N° alertas": 3})])
        data = r.json()["matched"][0]["result_data"]
        self.assertNotIn("n_alertas", data)
        self.assertNotIn("N° alertas", data)

    def test_the_first_measurement_falls_back_to_its_own_mean(self):
        # Sin basal previo y sin el checkbox, `cff_basal` cae a la propia media
        # y Δ% queda en 0 — la fórmula documentada en el seed.
        r = self._post([self._row("Elías Rojas")])
        data = r.json()["matched"][0]["result_data"]
        self.assertEqual(data["cff_basal"], data["cff_mean"])
        self.assertEqual(data["delta_basal_pct"], 0.0)

    def test_committing_writes_one_result_per_player_on_the_given_date(self):
        r = self._post(
            [self._row("Charles Aranguiz"), self._row("Elías Rojas")],
            dry_run=False,
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["created_results"], 2)

        rows = ExamResult.objects.filter(template=self.template)
        self.assertEqual(rows.count(), 2)
        for er in rows:
            self.assertEqual(er.recorded_at.date().isoformat(), "2026-08-20")

    # ── los dos bugs que esto destapó ───────────────────────────────────
    def test_a_naive_date_does_not_blow_up(self):
        """`BulkIngestForm` manda "…T12:00:00", sin offset.

        Eso llegaba naive al `ExamResult`, y el `post_save` de goals llamaba
        `timezone.localtime()` → ValueError → 500. La UI de carga masiva nunca
        se había ejercido completa, así que nadie lo había visto.

        ⚠️ `PR=2` no es un valor cualquiera: cae en la banda "Bajo", que tiene
        `alert: True`. `localtime()` se llama SÓLO cuando una regla dispara, así
        que con valores normales este test pasa sin ejercer el bug. Verificado
        quitando `make_aware`: con PR=7 pasa igual, con PR=2 falla.
        """
        r = self._post([self._row("Charles Aranguiz", PR=2)], dry_run=False,
                       recorded_at="2026-08-20T12:00:00")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(ExamResult.objects.filter(template=self.template).count(), 1)

    def test_the_stored_date_keeps_the_day_the_user_chose(self):
        # `make_aware` y no `replace(tzinfo=…)`: el string es hora local, y
        # reetiquetarlo como UTC corre el día en una toma nocturna.
        self._post([self._row("Charles Aranguiz", PR=2)], dry_run=False,
                   recorded_at="2026-08-20T22:00:00")
        er = ExamResult.objects.get(template=self.template)
        from django.utils import timezone

        self.assertEqual(
            timezone.localtime(er.recorded_at).date().isoformat(), "2026-08-20")

    def test_a_failure_mid_commit_leaves_nothing_behind(self):
        """El commit es atómico.

        Cada `create` dispara señales, así que una fila que falla dejaba las
        anteriores escritas y devolvía 500: el usuario veía un error con la
        mitad de sus datos adentro y sin saber cuál mitad.
        """
        rows = [self._row("Charles Aranguiz"), self._row("Elías Rojas")]
        calls = {"n": 0}
        real = ExamResult.objects.create

        def explode(*a, **kw):
            calls["n"] += 1
            if calls["n"] == 2:
                raise RuntimeError("falla simulada en la segunda fila")
            return real(*a, **kw)

        with mock.patch.object(ExamResult.objects, "create", side_effect=explode):
            with self.assertRaises(RuntimeError):
                self._post(rows, dry_run=False)

        self.assertEqual(ExamResult.objects.filter(template=self.template).count(), 0)


class BlankTemplateDownloadTests(FatigaCentralUploadTests):
    """La plantilla en blanco se genera desde el mapping, no es un archivo fijo."""

    def _download(self):
        return self.client.get(
            f"/api/templates/{self.template.id}/bulk-template.xlsx")

    def test_it_returns_an_xlsx(self):
        r = self._download()
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("plantilla-fatiga_central.xlsx", r["Content-Disposition"])

    def test_the_columns_follow_the_exams_field_order(self):
        """El orden NO puede venir de `field_map`.

        `input_config` es un JSONField y Postgres lo guarda como `jsonb`, que
        reordena las claves: esta plantilla vuelve como `EA, PR, I1, I2, I3`.
        El orden sale de los campos del examen, que es el orden en que el
        equipo mide.
        """
        r = self._download()
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb.active
        got = [ws.cell(1, j).value for j in range(1, ws.max_column + 1)]
        self.assertEqual(got, ["Jugador", "I1 (Hz)", "I2 (Hz)", "I3 (Hz)", "PR", "EA"])

        # …y para que quede claro que no es casualidad: así vuelve de la base.
        raw = list(
            self.template.input_config["column_mapping"]["field_map"].keys())
        self.assertNotEqual(raw, got[1:])

    def test_it_carries_no_data_rows(self):
        wb = openpyxl.load_workbook(io.BytesIO(self._download().content))
        self.assertEqual(wb.active.max_row, 1)

    def test_the_ignored_columns_are_not_offered(self):
        # Ponerlas invitaría a llenarlas para nada: el parser no las lee.
        wb = openpyxl.load_workbook(io.BytesIO(self._download().content))
        ws = wb.active
        got = {ws.cell(1, j).value for j in range(1, ws.max_column + 1)}
        for ignored in ("CFF mean", "CFF basal", "Δ% vs basal",
                        "Var % intra-sesión", "N° alertas"):
            self.assertNotIn(ignored, got, ignored)

    def test_a_file_built_from_the_downloaded_headers_ingests(self):
        """El cierre del círculo: lo que se descarga es lo que se puede subir."""
        wb = openpyxl.load_workbook(io.BytesIO(self._download().content))
        ws = wb.active
        headers = [ws.cell(1, j).value for j in range(1, ws.max_column + 1)]

        out = openpyxl.Workbook()
        os_ = out.active
        for j, h in enumerate(headers, 1):
            os_.cell(row=1, column=j, value=h)
        for j, v in enumerate(["Charles Aranguiz", 41.2, 40.8, 41.6, 7, 8], 1):
            os_.cell(row=2, column=j, value=v)
        buf = io.BytesIO()
        out.save(buf)

        r = self.client.post("/api/results/bulk", {
            "file": io.BytesIO(buf.getvalue()),
            "template_id": str(self.template.id),
            "category_id": str(self.cat.id),
            "recorded_at": "2026-08-20T12:00:00",
            "dry_run": "false",
        })
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["created_results"], 1)

    def test_a_template_without_bulk_mode_is_refused(self):
        self.template.input_config = {
            **self.template.input_config,
            "input_modes": ["single"],
        }
        self.template.save(update_fields=["input_config"])
        self.assertEqual(self._download().status_code, 400)
