"""One reading surface for the Formativo workbooks, xlsx or live Google Sheet.

The club keeps these documents in Google Sheets and hands out .xlsx exports of
them. The importers were written against the exports; pointing them at the live
document is what makes a scheduled sync possible. Both sources land here so the
parsers never learn which one they are reading.

Three differences the exports hide, each of which would corrupt data silently:

1. **Tab names are truncated in an export.** Excel caps a sheet title at 31
   characters, so the live `FORMATO CONDICIONAL 15-16 y 18-20` arrives as
   `FORMATO CONDICIONAL 15-16 y 18-`. Callers match by prefix (`match_sheet`).
2. **Rendered numbers are lossy.** The same `DT (m)` cell reads `4.159` as
   text and `4158.9` raw, because the dot is a THOUSANDS separator. Reading the
   rendered value divides every distance by a thousand and nothing complains,
   so the Sheets side always requests `UNFORMATTED_VALUE`.
3. **Rendered dates are ambiguous.** They come back as `8/01/2025`, which is
   the day/month coin-flip that already cost this project two birth dates.
   Unformatted they are serial numbers, which are exact.

Serial numbers are converted **only in date columns**, never blanket: a
`DURACIÓN (m)` of 37 is a valid serial too, and would quietly become
1900-02-05.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from typing import Any, Protocol


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c)).upper()
    return " ".join(re.sub(r"[^A-Z0-9 ]", " ", text).split())


def match_sheet(nombres: list[str], buscado: str) -> str | None:
    """The real tab title for a name that may be its 31-char truncation.

    Exact first, then prefix in either direction — the export truncates, so the
    stored constant can be shorter than the live title or the other way round
    if someone later shortens a tab.
    """
    if buscado in nombres:
        return buscado
    clave = _norm(buscado)
    for n in nombres:
        if _norm(n) == clave:
            return n
    for n in nombres:
        a, b = _norm(n), clave
        if a.startswith(b) or b.startswith(a):
            return n
    return None


class Fuente(Protocol):
    """A workbook: named sheets of positional rows, dates already real dates."""

    def hojas(self) -> list[str]: ...
    def filas(self, hoja: str) -> list[list[Any]]: ...


class FuenteXlsx:
    """An .xlsx export. openpyxl already hands back real datetimes."""

    def __init__(self, path: str):
        self.path = path
        self._book = None

    def _abrir(self):
        if self._book is None:
            import openpyxl

            self._book = openpyxl.load_workbook(self.path, data_only=True)
        return self._book

    def hojas(self) -> list[str]:
        return list(self._abrir().sheetnames)

    def filas(self, hoja: str) -> list[list[Any]]:
        real = match_sheet(self.hojas(), hoja)
        if real is None:
            return []
        return [list(r) for r in self._abrir()[real].iter_rows(values_only=True)]

    def cerrar(self):
        if self._book is not None:
            self._book.close()
            self._book = None


# A column holds dates when its header says so. `FECHA DE NACIMIENTO` and the
# bare `FECHA ` (with the club's trailing space) both match; `FECHA (S)` in the
# 1000 m sheet does not exist, but `TIEMPO (S)` would never match either.
_RE_FECHA = re.compile(r"^FECHA\b")


class FuenteSheets:
    """The live Google Sheet. Values come unformatted; dates come as serials."""

    def __init__(self, sheet_id: str, *, creds_file: str = "",
                 creds_json: str = ""):
        self.sheet_id = sheet_id
        self.creds = {"creds_file": creds_file, "creds_json": creds_json}
        self._hojas: list[str] | None = None
        self._cache: dict[str, list[list[Any]]] = {}

    def hojas(self) -> list[str]:
        from integrations.google_sheets import list_worksheets

        if self._hojas is None:
            self._hojas = list_worksheets(self.sheet_id, **self.creds)
        return self._hojas

    def filas(self, hoja: str) -> list[list[Any]]:
        from integrations.google_sheets import fetch_values, serial_to_date

        real = match_sheet(self.hojas(), hoja)
        if real is None:
            return []
        if real in self._cache:
            return self._cache[real]
        grid = fetch_values(self.sheet_id, real, **self.creds)
        if grid:
            # Which columns are dates is read from the header, once.
            cols = {i for i, h in enumerate(grid[0])
                    if _RE_FECHA.match(_norm(h))}
            for fila in grid[1:]:
                for i in cols:
                    if i < len(fila):
                        d = serial_to_date(fila[i])
                        if d is not None:
                            fila[i] = datetime.combine(d, datetime.min.time())
        self._cache[real] = grid
        return grid

    def cerrar(self):
        self._cache.clear()


def abrir(origen: str, *, creds_file: str = "", creds_json: str = "") -> Fuente:
    """`origen` is a filesystem path or a Google Sheets document id."""
    if origen.lower().endswith((".xlsx", ".xlsm")):
        return FuenteXlsx(origen)
    return FuenteSheets(origen, creds_file=creds_file, creds_json=creds_json)
