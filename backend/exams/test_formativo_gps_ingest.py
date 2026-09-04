"""Guards for `exams.formativo_gps_ingest`.

The bug this file exists for: the club bakes its thresholds into the column
headers (`HSR (m) >20km/h`, `AC (#) >3ms2`, `SPRINT (m) >25km/h`), so an exact
header map is pinned to those numbers. The first version was, and five of the
twelve metrics went unmapped — `acc`, `dec`, `hsr`, `sprint_dist`, `sprints`,
most of the point of a GPS import. It reported 15.352 rows and "0 filas sin
métricas", because the other seven columns did map.

So the mapping is by pattern, and a sheet that fails to yield an expected
metric says so instead of loading partial rows quietly.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import TestCase
from django.utils import timezone

from core.management.commands.backfill_cohorts import BRACKETS
from core.models import Bracket, Category, Club, Department, Player
from events.models import Event
from exams import formativo_gps_ingest as ing
from exams.models import ExamResult, ExamTemplate

# The club's real headers, thresholds included.
CAB = ["FECHA", "JUGADOR", "FECHA DE NACIMIENTO", "EDAD", "CATEGORÍA",
       "POSICIÓN", "CÓDIGO", "MICRO", "DURACIÓN (m)", "DT (m)", "MM",
       "AC (#) >3ms2", "DEC (#) >3ms2", "VMÁX (km/h)", "PL (UA)",
       "HSR (m) >20km/h", "SPRINT (m) >25km/h", "SPRINT (#) >25km/h",
       "Heart Rate - AVG HR  (% of player max HR)",
       "Heart Rate - Max Heart Rate (BPM)", "OBSERVACIÓN", "LOCALIA",
       "CALIDAD OPONENTE", "RESULTADO"]

D = datetime(2026, 3, 11)


def fila(nombre="CESAR PEREZ GALLEGOS", dia=D, codigo="MD-4", *,
         localia=None, calidad=None, resultado=None, hr=(None, None),
         vmax=21.8, dur=36.8):
    return [dia, nombre, datetime(2011, 6, 23), 15, "U15", "EXTREMO", codigo,
            1.0, dur, 2856.4, 77.6, 27.0, 10.0, vmax, 42.6, 262.4, 118.0, 3.0,
            hr[0], hr[1], "", localia, calidad, resultado]


def escribir(ruta: Path, hojas: dict[str, list[list]]) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for nombre, filas in hojas.items():
        ws = book.create_sheet(nombre)
        ws.append(CAB)
        for f in filas:
            ws.append(f)
    book.save(ruta)


class FormativoGpsTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.club = Club.objects.create(name="Universidad de Chile")
        cls.dept = Department.objects.create(club=cls.club, name="Físico",
                                             slug="fisico")
        for order, (code, name, age) in enumerate(BRACKETS):
            Bracket.objects.create(code=code, name=name, age=age, order=order)
        cls.cat = Category.objects.create(club=cls.club, name="Serie 2011",
                                          cohort_year=2011)
        cls.player = Player.objects.create(
            category=cls.cat, first_name="Cesar", last_name="Perez",
            second_last_name="Gallegos", date_of_birth=date(2011, 6, 23))
        campos = [
            {"key": "fecha", "label": "Fecha", "type": "date"},
            {"key": "sesion", "label": "Sesión", "type": "text"},
            {"key": "tot_dur", "label": "Duración", "type": "number", "unit": "min"},
            {"key": "tot_dist", "label": "DT", "type": "number", "unit": "m"},
            {"key": "mpm", "label": "MM", "type": "number", "unit": "m/min"},
            {"key": "acc", "label": "Acc", "type": "number"},
            {"key": "dec", "label": "Dec", "type": "number"},
            {"key": "acc_dec", "label": "Acc+Dec", "type": "number"},
            {"key": "max_vel", "label": "Vmáx", "type": "number", "unit": "km/h"},
            {"key": "player_load", "label": "PL", "type": "number"},
            {"key": "hsr", "label": "HSR", "type": "number", "unit": "m"},
            {"key": "sprint_dist", "label": "Sprint m", "type": "number", "unit": "m"},
            {"key": "sprints", "label": "Sprints", "type": "number"},
            {"key": "avg_hr_pct", "label": "FC media", "type": "number",
             "unit": "%", "min": 30, "max": 100},
            {"key": "max_hr_bpm", "label": "FC máx", "type": "number",
             "unit": "bpm", "min": 100, "max": 230},
        ]
        for slug, nombre in (("gps_sesion", "GPS sesión"),
                             ("gps_partido", "GPS partido")):
            extra = ([{"key": "tipo_sesion", "label": "Tipo",
                       "type": "categorical",
                       "options": ["entrenamiento", "reintegro"]}]
                     if slug == "gps_sesion" else [])
            t = ExamTemplate.objects.create(
                department=cls.dept, name=nombre, slug=slug,
                config_schema={"fields": campos + extra})
            t.applicable_categories.set([cls.cat])

    def correr(self, hojas, **kw):
        with TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "gps.xlsx"
            escribir(ruta, hojas)
            return ing.run(str(ruta), self.club, **kw)

    # ── el bug que motivó el archivo ────────────────────────────────────
    def test_mapea_las_metricas_aunque_el_umbral_este_en_el_encabezado(self):
        """`HSR (m) >20km/h` tiene que llegar a `hsr`.

        Con un mapa exacto, cambiar 20 por 21 en la planilla vuelve a dejar la
        métrica afuera sin que nada falle.
        """
        self.correr({"U15": [fila()]}, commit=True)
        d = ExamResult.objects.get(template__slug="gps_sesion").result_data
        for key, esperado in (("tot_dur", 36.8), ("tot_dist", 2856.4),
                              ("mpm", 77.6), ("acc", 27.0), ("dec", 10.0),
                              ("max_vel", 21.8), ("player_load", 42.6),
                              ("hsr", 262.4)):
            self.assertEqual(d.get(key), esperado, key)

    def test_distingue_sprint_en_metros_de_sprint_en_cuenta(self):
        # Los dos encabezados normalizan casi igual; el orden los separa.
        self.correr({"U15": [fila()]}, commit=True)
        d = ExamResult.objects.get(template__slug="gps_sesion").result_data
        self.assertEqual(d["sprint_dist"], 118.0)
        self.assertEqual(d["sprints"], 3.0)

    def test_avisa_cuando_una_metrica_esperada_no_mapea(self):
        """El silencio es peor que el error.

        Si la hoja cambia de forma, la corrida tiene que decirlo en vez de
        cargar filas a medias.
        """
        with TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "gps.xlsx"
            sin_hsr = [c for c in CAB if not c.startswith("HSR")]
            import openpyxl
            book = openpyxl.Workbook()
            book.remove(book.active)
            ws = book.create_sheet("U15")
            ws.append(sin_hsr)
            book.save(ruta)
            _, faltantes, _, _ = ing.parse_workbook(str(ruta))
        self.assertIn("hsr", faltantes.get("U15", []))

    def test_avisa_cuando_aparece_una_columna_nueva(self):
        """El riesgo espejo, y el silencioso.

        La planilla la mantienen a mano. Que falte una métrica ya avisa; que
        AGREGUEN una columna no avisaba nada — simplemente no llegaba, y la
        corrida seguía reportando una carga completa.
        """
        with TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "gps.xlsx"
            import openpyxl
            book = openpyxl.Workbook()
            book.remove(book.active)
            ws = book.create_sheet("U15")
            ws.append(CAB + ["HMLD (m)"])
            ws.append(fila() + [1022.4])
            book.save(ruta)
            _, faltantes, desconocidas, _ = ing.parse_workbook(str(ruta))
        self.assertEqual(faltantes, {}, "no falta ninguna métrica esperada")
        self.assertEqual(desconocidas.get("U15"), ["HMLD (m)"])

    def test_no_avisa_por_las_columnas_que_se_ignoran_a_proposito(self):
        # Identidad, microciclo, observación y las tres marcas de partido no
        # son métricas y no deben ensuciar el aviso.
        with TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "gps.xlsx"
            escribir(ruta, {"U15": [fila()]})
            _, _, desconocidas, _ = ing.parse_workbook(str(ruta))
        self.assertEqual(desconocidas, {})

    # ── partido o sesión ────────────────────────────────────────────────
    def test_las_marcas_del_club_definen_que_es_partido(self):
        rep = self.correr({"U15": [
            fila(codigo="MD", localia="LOCAL", calidad="ALTA", resultado="2-1"),
            fila(dia=datetime(2026, 3, 14), codigo="MD-2"),
        ]}, commit=True)
        self.assertEqual(rep.partidos, 1)
        self.assertEqual(rep.sesiones, 1)
        self.assertTrue(ExamResult.objects.filter(template__slug="gps_partido").exists())

    def test_dia_de_partido_sin_marcas_va_a_sesion(self):
        """2174 filas del archivo son así, y no parecen partidos.

        Median 47,0 min y 5180 m, contra 82,2 min y 8214 m de las marcadas.
        Y sin marcas no hay rival ni resultado, así que un `gps_partido` sería
        un registro de partido sin partido.
        """
        rep = self.correr({"U15": [fila(codigo="MD")]}, commit=True)
        self.assertEqual(rep.partidos, 0)
        self.assertEqual(rep.md_sin_marca, 1)
        r = ExamResult.objects.get()
        self.assertEqual(r.template.slug, "gps_sesion")
        self.assertEqual(r.result_data["tipo_sesion"], "entrenamiento")

    def test_una_sola_marca_alcanza(self):
        # 577 filas del archivo traen localía y resultado pero no calidad.
        rep = self.correr({"U15": [
            fila(codigo="MD", localia="VISITA", resultado="0-3")]}, commit=True)
        self.assertEqual(rep.partidos, 1)

    def test_vincula_el_partido_al_evento_de_ese_dia(self):
        ev = Event.objects.create(
            club=self.club, category=self.cat, department=self.dept,
            event_type="match", title="vs Rival",
            starts_at=timezone.make_aware(D))
        self.correr({"U15": [
            fila(codigo="MD", localia="LOCAL", resultado="1-0")]}, commit=True)
        self.assertEqual(ExamResult.objects.get().event, ev)

    def test_un_partido_sin_evento_en_SLAB_igual_se_importa(self):
        # El archivo arranca en 2024 y los eventos COMET en 2025-01-23: toda
        # la primera temporada no tiene evento al que colgarse.
        rep = self.correr({"U15": [
            fila(codigo="MD", localia="LOCAL", resultado="1-0")]}, commit=True)
        self.assertEqual(rep.creados, 1)
        self.assertIsNone(ExamResult.objects.get().event)
        self.assertEqual(rep.con_evento, 0)

    # ── hojas que no van ────────────────────────────────────────────────
    def test_excluye_las_hojas_de_otros_equipos_y_los_bloques(self):
        """`U17WC`/`U20 WC`/`WC CLUBES` traen jugadores de Panamá, Guatemala,
        Al Ahly e Inter Miami. `Pasing Indi.` parte una sesión en bloques de
        15 minutos, así que importarla contaría el mismo trabajo varias veces.
        """
        rep = self.correr({"U15": [fila()], "U17WC": [fila()],
                           "WC CLUBES": [fila()], "Pasing Indi.": [fila()],
                           "GPS PARTIDOS": [fila()]}, commit=True)
        self.assertEqual(rep.creados, 1)
        self.assertEqual(dict(rep.por_hoja), {"U15": 1})

    # ── decisiones de campo ─────────────────────────────────────────────
    def test_acc_dec_queda_vacio(self):
        """Catapult tampoco lo llena (no está en su SLUG_MAP).

        Sumar AC+DEC acá haría del formativo el único origen que lo trae, que
        es la misma inconsistencia al revés.
        """
        self.correr({"U15": [fila()]}, commit=True)
        self.assertNotIn("acc_dec", ExamResult.objects.get().result_data)

    def test_un_pulso_en_cero_no_es_una_medicion(self):
        rep = self.correr({"U15": [fila(hr=(0.0, 0.0))]}, commit=True)
        d = ExamResult.objects.get().result_data
        self.assertNotIn("avg_hr_pct", d)
        self.assertNotIn("max_hr_bpm", d)
        self.assertEqual(len(rep.fuera_de_rango), 1)
        self.assertEqual(rep.creados, 1, "la fila igual se carga")

    def test_guarda_el_pulso_cuando_viene(self):
        self.correr({"U15": [fila(hr=(82.0, 191.0))]}, commit=True)
        d = ExamResult.objects.get().result_data
        self.assertEqual(d["avg_hr_pct"], 82.0)
        self.assertEqual(d["max_hr_bpm"], 191.0)

    def test_el_codigo_de_microciclo_va_como_procedencia_no_como_metrica(self):
        # `exams.microcycle` deriva MD-n de las fechas de partido del club.
        self.correr({"U15": [fila(codigo="MD-4")]}, commit=True)
        d = ExamResult.objects.get().result_data
        self.assertEqual(d["origen_codigo"], "MD-4")
        self.assertIn("MD-4", d["sesion"])

    # ── procedencia e idempotencia ──────────────────────────────────────
    def test_escribe_la_procedencia(self):
        self.correr({"U15": [fila()]}, commit=True)
        d = ExamResult.objects.get().result_data
        self.assertEqual(d["origen"], ing.ORIGEN)
        self.assertEqual(d["origen_hoja"], "U15")
        self.assertNotEqual(d["origen"], "planilla_club_formativo",
                            "distinto del import físico, para poder filtrar")

    def test_correrlo_dos_veces_no_duplica(self):
        hojas = {"U15": [fila()]}
        self.correr(hojas, commit=True)
        rep = self.correr(hojas, commit=True)
        self.assertEqual(rep.creados, 0)
        self.assertEqual(rep.ya_existian, 1)
        self.assertEqual(ExamResult.objects.count(), 1)

    def test_la_misma_fecha_con_codigos_distintos_son_dos_filas(self):
        # Un jugador puede tener sesión y partido el mismo día.
        rep = self.correr({"U15": [
            fila(codigo="MD-1"),
            fila(codigo="MD", localia="LOCAL", resultado="1-0"),
        ]}, commit=True)
        self.assertEqual(rep.creados, 2)

    def test_sin_commit_no_escribe(self):
        rep = self.correr({"U15": [fila()]})
        self.assertEqual(rep.creados, 1)
        self.assertEqual(ExamResult.objects.count(), 0)

    def test_una_fila_sin_fecha_se_reporta(self):
        rep = self.correr({"U15": [fila(dia=None)]}, commit=True)
        self.assertEqual(rep.creados, 0)
        self.assertEqual(rep.sin_fecha, 1)
