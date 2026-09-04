"""Guards for `exams.formativo_sources` — the xlsx/Sheets reading surface.

Every failure mode here is silent and numeric, which is the worst combination:

* the rendered `DT (m)` cell reads `4.159` where the value is `4158.9`, because
  the dot is a THOUSANDS separator — read the rendered text and every distance
  is off by a thousand with nothing to notice;
* rendered dates are `8/01/2025`, the day/month coin-flip that already cost
  this project two birth dates;
* an .xlsx export truncates a tab title to 31 characters, so the constant
  `FORMATO CONDICIONAL 15-16 y 18-` does not match the live
  `FORMATO CONDICIONAL 15-16 y 18-20` and the sheet is skipped in silence.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from django.test import SimpleTestCase

from exams.formativo_sources import (FuenteSheets, FuenteXlsx, abrir,
                                     match_sheet)
from integrations.google_sheets import SERIAL_EPOCH, serial_to_date


class SerialADateTests(SimpleTestCase):
    """The unformatted read hands back serial numbers; these are exact."""

    def test_convierte_un_serial_conocido(self):
        # Verificado contra la hoja viva: 45665 es 2025-01-08, y el nacimiento
        # 38881 es 2006-06-13 (la celda formateada dice 13/06/2006).
        self.assertEqual(serial_to_date(45665), date(2025, 1, 8))
        self.assertEqual(serial_to_date(38881), date(2006, 6, 13))

    def test_la_epoca_es_1899_12_30(self):
        # Sheets hereda de Excel el bug del año 1900 bisiesto, que corre el
        # cero un día.
        self.assertEqual(SERIAL_EPOCH, date(1899, 12, 30))
        self.assertEqual(serial_to_date(1), date(1899, 12, 31))

    def test_lo_que_no_es_un_serial_devuelve_None(self):
        for v in ("2025-01-08", None, "", True, False, 0, 0.5, 999999):
            self.assertIsNone(serial_to_date(v), repr(v))


class MatchSheetTests(SimpleTestCase):
    VIVAS = ["JUGADORES", "FORMATO CONDICIONAL 15-16 y 18-20",
             "FORMATO CONDICIONAL 13 -14 y 11-12", "PRESS DE BANCO", "CARRERAS"]

    def test_resuelve_el_nombre_truncado_del_export(self):
        """El caso real: Excel corta el título a 31 caracteres.

        Sin esto, las dos hojas de FORMATO CONDICIONAL se saltan y
        `seed_formativo_bands` no siembra ninguna banda — sin error.
        """
        self.assertEqual(
            match_sheet(self.VIVAS, "FORMATO CONDICIONAL 15-16 y 18-"),
            "FORMATO CONDICIONAL 15-16 y 18-20")
        self.assertEqual(
            match_sheet(self.VIVAS, "FORMATO CONDICIONAL 13 -14 y 11"),
            "FORMATO CONDICIONAL 13 -14 y 11-12")

    def test_el_nombre_exacto_gana(self):
        self.assertEqual(match_sheet(self.VIVAS, "CARRERAS"), "CARRERAS")

    def test_no_confunde_dos_hojas_distintas(self):
        # `CARRERAS` no debe resolver a `PRESS DE BANCO` ni al revés.
        self.assertEqual(match_sheet(self.VIVAS, "PRESS DE BANCO"),
                         "PRESS DE BANCO")

    def test_una_hoja_que_no_existe_devuelve_None(self):
        self.assertIsNone(match_sheet(self.VIVAS, "HOJA INVENTADA"))


class FuenteXlsxTests(SimpleTestCase):
    def _libro(self, tmp: str) -> str:
        import openpyxl

        book = openpyxl.Workbook()
        book.remove(book.active)
        ws = book.create_sheet("FORMATO CONDICIONAL 15-16 y 18-20")
        ws.append(["A", "B"])
        ws.append([1, 2])
        ws2 = book.create_sheet("CARRERAS")
        ws2.append(["FECHA", "DT (m)"])
        ws2.append([datetime(2026, 3, 11), 4158.9])
        ruta = str(Path(tmp) / "x.xlsx")
        book.save(ruta)
        return ruta

    def test_devuelve_la_grilla_posicional_con_fechas_reales(self):
        with TemporaryDirectory() as tmp:
            f = FuenteXlsx(self._libro(tmp))
            grid = f.filas("CARRERAS")
            self.assertEqual(grid[0], ["FECHA", "DT (m)"])
            self.assertEqual(grid[1][0], datetime(2026, 3, 11))
            self.assertEqual(grid[1][1], 4158.9)
            f.cerrar()

    def test_tambien_resuelve_por_prefijo(self):
        # Un .xlsx puede tener el título completo o el truncado; los dos valen.
        with TemporaryDirectory() as tmp:
            f = FuenteXlsx(self._libro(tmp))
            self.assertTrue(f.filas("FORMATO CONDICIONAL 15-16 y 18-"))
            f.cerrar()

    def test_una_hoja_inexistente_devuelve_vacio_sin_reventar(self):
        with TemporaryDirectory() as tmp:
            f = FuenteXlsx(self._libro(tmp))
            self.assertEqual(f.filas("NO EXISTE"), [])
            f.cerrar()


class FuenteSheetsTests(SimpleTestCase):
    """Con la API simulada: lo que importa es la coerción, no la red."""

    GRID = [
        ["FECHA ", "JUGADOR", "FECHA DE NACIMIENTO", "DURACIÓN (m)", "DT (m)"],
        [45665, "ALONSO VILLEGAS MUNITA", 38881, 36.6, 4158.9],
        [45665, "OTRO JUGADOR", "", 37, 4130.2],
    ]

    def _fuente(self):
        f = FuenteSheets("id-falso", creds_file="/dev/null")
        f._hojas = ["U20"]
        return f

    def test_convierte_a_fecha_solo_las_columnas_de_fecha(self):
        """La coerción es por columna, nunca a ciegas.

        Una `DURACIÓN (m)` de 37 también es un serial válido, y a ciegas se
        volvería 1900-02-05.
        """
        f = self._fuente()
        with _fake_fetch(self.GRID):
            grid = f.filas("U20")
        self.assertEqual(grid[1][0], datetime(2025, 1, 8))     # FECHA
        self.assertEqual(grid[1][2], datetime(2006, 6, 13))    # NACIMIENTO
        self.assertEqual(grid[1][3], 36.6, "la duración NO es una fecha")
        self.assertEqual(grid[1][4], 4158.9, "la distancia NO es una fecha")

    def test_una_celda_de_fecha_vacia_queda_como_esta(self):
        f = self._fuente()
        with _fake_fetch(self.GRID):
            grid = f.filas("U20")
        self.assertEqual(grid[2][2], "")

    def test_un_entero_en_columna_de_fecha_sigue_siendo_entero_si_no_es_serial(self):
        grid_raro = [["FECHA "], [0]]
        f = self._fuente()
        with _fake_fetch(grid_raro):
            out = f.filas("U20")
        self.assertEqual(out[1][0], 0, "0 no es una fecha")

    def test_cachea_la_hoja_para_no_pedirla_dos_veces(self):
        # Cada lectura es una llamada de red; el parser recorre la hoja varias
        # veces buscando bloques.
        f = self._fuente()
        with _fake_fetch(self.GRID) as calls:
            f.filas("U20")
            f.filas("U20")
        self.assertEqual(calls["n"], 1)


class AbrirTests(SimpleTestCase):
    def test_una_ruta_xlsx_da_la_fuente_de_archivo(self):
        self.assertIsInstance(abrir("/tmp/x.xlsx"), FuenteXlsx)
        self.assertIsInstance(abrir("/tmp/X.XLSM"), FuenteXlsx)

    def test_cualquier_otra_cosa_se_trata_como_id_de_hoja(self):
        self.assertIsInstance(
            abrir("1WHlhb-K1Pbk1-ttkkCzdIMT69RCfYzXhDyzFXMI4_vo"), FuenteSheets)


class _fake_fetch:
    """Reemplaza `fetch_values` por una grilla fija y cuenta las llamadas."""

    def __init__(self, grid):
        self.grid = grid
        self.calls = {"n": 0}

    def __enter__(self):
        import integrations.google_sheets as gs

        self._orig = gs.fetch_values

        def fake(sheet_id, worksheet, **kw):
            self.calls["n"] += 1
            return [list(r) for r in self.grid]

        gs.fetch_values = fake
        return self.calls

    def __exit__(self, *exc):
        import integrations.google_sheets as gs

        gs.fetch_values = self._orig
        return False
