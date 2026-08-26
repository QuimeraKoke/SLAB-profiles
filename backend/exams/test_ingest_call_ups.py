"""Every importer must resolve CALLED-UP players, not just the home roster.

The bug, found while counting phase-4 sites: `bulk_ingest` widened its player
index to home ∪ active call-ups, but `gps_session_ingest`, `wellness_ingest` and
`catapult_sync` never did. A called-up player appears in the provider's export
for the squad he actually trained with, so his rows landed in `unmatched` and
were dropped. On Universidad de Chile 2026, 26% of youth appearances are loans.

Two halves have to move together, which is the part worth testing:

  * the NAME INDEX, or the row never resolves — data lost;
  * the DEDUP lookup, or the resolved row is written again on every run —
    data duplicated.

Fixing only the first turns a silent loss into a silent duplication.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
from django.test import TestCase

from core.models import Category, Club, Department, Player, PlayerCallUp
from exams import gps_session_ingest
from exams.gps_session import HEADER_TO_KEY, PLAYER_COL, SESSION_COL
from exams.models import ExamResult, ExamTemplate


def _workbook(rows):
    """A minimal per-session GPS export: Players, Sessions + the metric columns."""
    headers = [PLAYER_COL, SESSION_COL, *HEADER_TO_KEY]
    wb = openpyxl.Workbook()
    ws = wb.active
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h)
    for i, (player, session) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=player)
        ws.cell(row=i, column=2, value=session)
        for j in range(3, len(headers) + 1):
            ws.cell(row=i, column=j, value=10)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class GpsSessionIngestCallUpTests(TestCase):
    def setUp(self):
        self.club = Club.objects.create(name="FC")
        self.dept = Department.objects.create(
            club=self.club, name="Preparación física", slug="fisico")
        self.host = Category.objects.create(club=self.club, name="SUB-18")
        self.home = Category.objects.create(club=self.club, name="SUB-16")

        self.local = Player.objects.create(
            category=self.host, first_name="Local", last_name="Uno")
        self.lent = Player.objects.create(
            category=self.home, first_name="Prestado", last_name="Dos")
        PlayerCallUp.objects.create(
            player=self.lent, category=self.host, active=True, since=date(2026, 3, 1),
        )
        self.template = ExamTemplate.objects.create(
            department=self.dept, name="GPS sesión", slug="gps_sesion",
            config_schema={"fields": []},
        )

    @staticmethod
    def _unmatched(report):
        """`unmatched` is a list of {code, rows} — flatten to the labels."""
        return {u["code"] for u in report["unmatched"]}

    def _run(self, rows, **kw):
        return gps_session_ingest.run(
            _workbook(rows), template=self.template, category=self.host,
            department=self.dept, mode="training", **kw,
        )

    # ------------------------------------------------------------------
    def test_a_called_up_players_row_resolves_instead_of_being_dropped(self):
        report = self._run([
            ("Local Uno", "Entrenamiento 12-3"),
            ("Prestado Dos", "Entrenamiento 12-3"),
        ])
        self.assertEqual(self._unmatched(report), set())
        self.assertEqual(report["created"], 2)

    def test_a_stranger_is_still_not_matched(self):
        # Widening to call-ups must not become "match anyone in the club": an
        # unrelated player has to stay unmatched so the queue still catches
        # genuinely unknown names.
        Player.objects.create(
            category=self.home, first_name="Ajeno", last_name="Tres")
        report = self._run([("Ajeno Tres", "Entrenamiento 12-3")])
        self.assertEqual(report["created"], 0)
        self.assertIn("Ajeno Tres", self._unmatched(report))

    def test_the_call_up_row_is_written_to_the_database(self):
        self._run([("Prestado Dos", "Entrenamiento 12-3")], dry_run=False)
        self.assertTrue(
            ExamResult.objects.filter(template=self.template, player=self.lent).exists()
        )

    def test_re_running_does_not_duplicate_the_call_ups_row(self):
        # The second half of the bug. The existing row hangs off the player's
        # HOME category, so a dedup keyed on `player__category=host` would not
        # see it and would write it again — turning a silent loss into a silent
        # duplication.
        rows = [("Prestado Dos", "Entrenamiento 12-3")]
        self._run(rows, dry_run=False)
        report = self._run(rows, dry_run=False)

        self.assertEqual(report["created"], 0)
        self.assertEqual(report["skipped"], 1)
        self.assertEqual(
            ExamResult.objects.filter(
                template=self.template, player=self.lent).count(),
            1,
        )

    def test_an_inactive_call_up_stops_resolving(self):
        PlayerCallUp.objects.filter(player=self.lent).update(active=False)
        report = self._run([("Prestado Dos", "Entrenamiento 12-3")])
        self.assertIn("Prestado Dos", self._unmatched(report))


class WellnessMatcherCallUpTests(TestCase):
    def setUp(self):
        self.club = Club.objects.create(name="FC")
        self.host = Category.objects.create(club=self.club, name="SUB-18")
        self.home = Category.objects.create(club=self.club, name="SUB-16")
        self.lent = Player.objects.create(
            category=self.home, first_name="Prestado", last_name="Dos")
        PlayerCallUp.objects.create(
            player=self.lent, category=self.host, active=True, since=date(2026, 3, 1),
        )

    def test_a_called_up_players_check_in_resolves(self):
        # The wellness form is filled by whoever trained with the squad.
        from exams.wellness_ingest import _build_matcher

        match = _build_matcher(self.host)
        self.assertEqual(match("Prestado Dos"), self.lent)


class CatapultRosterCallUpTests(TestCase):
    def setUp(self):
        self.club = Club.objects.create(name="FC")
        self.host = Category.objects.create(club=self.club, name="SUB-18")
        self.home = Category.objects.create(club=self.club, name="SUB-16")
        self.lent = Player.objects.create(
            category=self.home, first_name="Prestado", last_name="Dos")
        PlayerCallUp.objects.create(
            player=self.lent, category=self.host, active=True, since=date(2026, 3, 1),
        )

    def test_a_called_up_player_is_in_the_catapult_index(self):
        # Catapult tags a session with the squad that trained, so the row
        # arrives under the host category.
        from exams.services.catapult_sync import _norm, _roster_index

        players, _by_dob, by_name = _roster_index(self.host)
        self.assertIn(self.lent, players)
        self.assertEqual(by_name[_norm("Prestado Dos")], self.lent)
